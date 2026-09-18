import io
from datetime import date, datetime

import openpyxl
from django.db.models import Q
from django.http import HttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.files.drive_service import upload_file, delete_file
from .models import SampleTrackerRecord, SampleTrackerFile
from .serializers import SampleTrackerRecordSerializer, SampleTrackerFileSerializer

TEMPLATE_COLUMNS = [
    'date', 'client_name', 'product_type', 'product_details',
    'sample_attempt_count', 'sample_no', 'sample_approved', 'feedback',
]

MAX_UPLOAD_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

_DATE_FORMATS = [
    '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%m/%d/%Y',
    '%b %Y', '%B %Y', '%b-%Y', '%B-%Y',
]


def _parse_date(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    raw = str(raw).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


class SampleTrackerRecordViewSet(viewsets.ModelViewSet):
    serializer_class = SampleTrackerRecordSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = SampleTrackerRecord.objects.all()
        q = self.request.query_params.get('q')
        if q:
            qs = qs.filter(
                Q(client_name__icontains=q) | Q(product_type__icontains=q) |
                Q(sample_no__icontains=q) | Q(product_details__icontains=q)
            )
        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete.'}, status=403)
        return super().destroy(request, *args, **kwargs)

    # ------------------------------------------------------------------
    # GET /api/sample-tracker/upload-template/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['get'], url_path='upload-template')
    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Final Formula Tracker'

        header_font = openpyxl.styles.Font(bold=True)
        header_fill = openpyxl.styles.PatternFill('solid', fgColor='FFF3CD')

        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill

        widths = [14, 22, 18, 30, 18, 14, 16, 30]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        hint_font = openpyxl.styles.Font(italic=True, color='888888')
        date_col = TEMPLATE_COLUMNS.index('date') + 1
        ws.cell(row=2, column=date_col, value='YYYY-MM-DD, DD-MM-YYYY, or month + year e.g. Feb 2026').font = hint_font

        sample = {
            'date': '2026-06-01',
            'client_name': 'Vedanshi',
            'product_type': 'Face Serum',
            'product_details': 'Vitamin C Serum, 30ml',
            'sample_attempt_count': 1,
            'sample_no': 'S-001',
            'sample_approved': 'Pending',
            'feedback': 'Awaiting client review',
        }
        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            ws.cell(row=3, column=col_idx, value=sample[col_name])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="final_formula_tracker_template.xlsx"'
        return response

    # ------------------------------------------------------------------
    # POST /api/sample-tracker/bulk-upload/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['post'], url_path='bulk-upload')
    def bulk_upload(self, request):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file provided. Send an Excel file as "file".'}, status=400)

        try:
            wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        except Exception:
            return Response({'detail': 'Could not read the file. Please upload a valid .xlsx file.'}, status=400)

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return Response({'detail': 'The file is empty.'}, status=400)

        header_row = [str(c).strip().lower() if c is not None else '' for c in rows[0]]

        def col(name):
            try:
                return header_row.index(name)
            except ValueError:
                return None

        idx = {name: col(name) for name in TEMPLATE_COLUMNS}

        def cell(row_data, name):
            i = idx[name]
            if i is None or i >= len(row_data):
                return None
            v = row_data[i]
            return str(v).strip() if v is not None else None

        def int_or_none(row_data, name):
            v = cell(row_data, name)
            if v is None:
                return None
            try:
                return int(float(v))
            except (ValueError, TypeError):
                return None

        created = []
        to_create = []

        for row_num, row_data in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == '' for v in row_data):
                continue

            date_raw = row_data[idx['date']] if idx['date'] is not None and idx['date'] < len(row_data) else None
            parsed_date = _parse_date(date_raw)

            to_create.append(SampleTrackerRecord(
                date=parsed_date,
                client_name=cell(row_data, 'client_name') or '',
                product_type=cell(row_data, 'product_type') or '',
                product_details=cell(row_data, 'product_details') or '',
                sample_attempt_count=int_or_none(row_data, 'sample_attempt_count'),
                sample_no=cell(row_data, 'sample_no') or '',
                sample_approved=cell(row_data, 'sample_approved') or '',
                feedback=cell(row_data, 'feedback') or '',
                created_by=request.user,
            ))
            entry = {'row': row_num, 'client_name': cell(row_data, 'client_name') or '—'}
            if date_raw and not parsed_date:
                entry['warning'] = f'Could not parse date "{date_raw}" — left blank'
            created.append(entry)

        SampleTrackerRecord.objects.bulk_create(to_create)
        wb.close()
        return Response({'created': created, 'skipped': []}, status=200)


class SampleTrackerFilesViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def list(self, request, sample_tracker_id=None):
        files = SampleTrackerFile.objects.filter(sample_tracker_record_id=sample_tracker_id).select_related('uploaded_by')
        return Response(SampleTrackerFileSerializer(files, many=True).data)

    def create(self, request, sample_tracker_id=None):
        try:
            record = SampleTrackerRecord.objects.get(pk=sample_tracker_id)
        except SampleTrackerRecord.DoesNotExist:
            raise NotFound('Sample tracker record not found')

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file uploaded'}, status=status.HTTP_400_BAD_REQUEST)
        if file_obj.size > MAX_UPLOAD_SIZE_BYTES:
            return Response({'detail': 'File exceeds the 10 MB upload limit.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            result = upload_file(
                file_bytes=file_obj.read(),
                filename=file_obj.name,
                mimetype=file_obj.content_type or 'application/octet-stream',
                client_name=record.client_name or 'Sample Tracker',
                subfolder='Sample Tracker',
            )
        except Exception as e:
            return Response({'detail': f'Drive upload failed: {e}'}, status=status.HTTP_502_BAD_GATEWAY)

        file_record = SampleTrackerFile.objects.create(
            sample_tracker_record=record,
            drive_file_id=result['drive_file_id'],
            drive_url=result['drive_url'],
            filename=file_obj.name,
            uploaded_by=request.user,
        )
        return Response(SampleTrackerFileSerializer(file_record).data, status=status.HTTP_201_CREATED)

    def destroy(self, request, pk=None, sample_tracker_id=None):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete files.'}, status=status.HTTP_403_FORBIDDEN)
        try:
            file_record = SampleTrackerFile.objects.get(pk=pk, sample_tracker_record_id=sample_tracker_id)
        except SampleTrackerFile.DoesNotExist:
            raise NotFound()
        delete_file(file_record.drive_file_id)
        file_record.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)
