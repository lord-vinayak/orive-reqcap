import io

import openpyxl
from django.db.models import Q
from django.http import HttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.files.drive_service import upload_file, delete_file
from .models import PackagingClientRecord, PackagingClientFile
from .serializers import PackagingClientRecordSerializer, PackagingClientFileSerializer

TEMPLATE_COLUMNS = [
    'client_name', 'packaging_name', 'size', 'glass_pet', 'client_moq', 'vendor_moq', 'cost_to_ss',
    'cost_to_client', 'vendor_name', 'poc', 'contact_details', 'poc2', 'cd2',
]

MAX_UPLOAD_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB


class PackagingClientRecordViewSet(viewsets.ModelViewSet):
    serializer_class = PackagingClientRecordSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = PackagingClientRecord.objects.all()
        q = self.request.query_params.get('q')
        if q:
            qs = qs.filter(
                Q(client_name__icontains=q) | Q(packaging_name__icontains=q) |
                Q(vendor_name__icontains=q) | Q(poc__icontains=q)
            )
        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete.'}, status=403)
        return super().destroy(request, *args, **kwargs)

    # ------------------------------------------------------------------
    # GET /api/packaging-clients/upload-template/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['get'], url_path='upload-template')
    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Packaging Clients'

        header_font = openpyxl.styles.Font(bold=True)
        header_fill = openpyxl.styles.PatternFill('solid', fgColor='FFF3CD')

        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill

        widths = [18, 22, 12, 12, 10, 10, 14, 14, 22, 16, 18, 16, 18]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        sample = {
            'client_name': 'Vedanshi',
            'packaging_name': 'PP Airless Bottle',
            'size': '50ml',
            'glass_pet': 'Pet',
            'client_moq': '1100',
            'vendor_moq': '1000',
            'cost_to_ss': '29rs',
            'cost_to_client': '35rs',
            'vendor_name': 'Indian Harness',
            'poc': 'Anjali',
            'contact_details': '99109 16449',
            'poc2': 'Radhika',
            'cd2': '88268 70902',
        }
        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            ws.cell(row=2, column=col_idx, value=sample[col_name])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="packaging_clients_template.xlsx"'
        return response

    # ------------------------------------------------------------------
    # POST /api/packaging-clients/bulk-upload/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['post'], url_path='bulk-upload')
    def bulk_upload(self, request):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file provided. Send an Excel file as "file".'}, status=400)

        try:
            wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        except Exception:
            return Response({'detail': 'Could not read the file. Please upload a valid .xlsx file.'}, status=400)

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return Response({'detail': 'The file is empty.'}, status=400)

        header_row = [str(c).strip().lower() if c is not None else '' for c in rows[0]]

        def col(name):
            try:
                return header_row.index(name)
            except ValueError:
                return None

        idx = {name: col(name) for name in TEMPLATE_COLUMNS}

        def cell(row_data, name):
            i = idx[name]
            if i is None or i >= len(row_data):
                return None
            v = row_data[i]
            return str(v).strip() if v is not None else None

        created = []
        to_create = []

        for row_num, row_data in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == '' for v in row_data):
                continue

            to_create.append(PackagingClientRecord(
                client_name=cell(row_data, 'client_name') or '',
                packaging_name=cell(row_data, 'packaging_name') or '',
                size=cell(row_data, 'size') or '',
                glass_pet=cell(row_data, 'glass_pet') or '',
                client_moq=cell(row_data, 'client_moq') or '',
                vendor_moq=cell(row_data, 'vendor_moq') or '',
                cost_to_ss=cell(row_data, 'cost_to_ss') or '',
                cost_to_client=cell(row_data, 'cost_to_client') or '',
                vendor_name=cell(row_data, 'vendor_name') or '',
                poc=cell(row_data, 'poc') or '',
                contact_details=cell(row_data, 'contact_details') or '',
                poc2=cell(row_data, 'poc2') or '',
                cd2=cell(row_data, 'cd2') or '',
                created_by=request.user,
            ))
            created.append({'row': row_num, 'client_name': cell(row_data, 'client_name') or '—'})

        PackagingClientRecord.objects.bulk_create(to_create)
        wb.close()
        return Response({'created': created, 'skipped': []}, status=200)


class PackagingClientFilesViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def list(self, request, packaging_client_id=None):
        files = PackagingClientFile.objects.filter(packaging_client_id=packaging_client_id).select_related('uploaded_by')
        return Response(PackagingClientFileSerializer(files, many=True).data)

    def create(self, request, packaging_client_id=None):
        try:
            record = PackagingClientRecord.objects.get(pk=packaging_client_id)
        except PackagingClientRecord.DoesNotExist:
            raise NotFound('Packaging client not found')

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file uploaded'}, status=status.HTTP_400_BAD_REQUEST)
        if file_obj.size > MAX_UPLOAD_SIZE_BYTES:
            return Response({'detail': 'File exceeds the 10 MB upload limit.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            result = upload_file(
                file_bytes=file_obj.read(),
                filename=file_obj.name,
                mimetype=file_obj.content_type or 'application/octet-stream',
                client_name=record.client_name or 'Packaging Clients',
                subfolder='Packaging Clients',
            )
        except Exception as e:
            return Response({'detail': f'Drive upload failed: {e}'}, status=status.HTTP_502_BAD_GATEWAY)

        file_record = PackagingClientFile.objects.create(
            packaging_client=record,
            drive_file_id=result['drive_file_id'],
            drive_url=result['drive_url'],
            filename=file_obj.name,
            uploaded_by=request.user,
        )
        return Response(PackagingClientFileSerializer(file_record).data, status=status.HTTP_201_CREATED)

    def destroy(self, request, pk=None, packaging_client_id=None):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete files.'}, status=status.HTTP_403_FORBIDDEN)
        try:
            file_record = PackagingClientFile.objects.get(pk=pk, packaging_client_id=packaging_client_id)
        except PackagingClientFile.DoesNotExist:
            raise NotFound()
        delete_file(file_record.drive_file_id)
        file_record.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)
