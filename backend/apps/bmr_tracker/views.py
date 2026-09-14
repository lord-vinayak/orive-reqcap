import io
from datetime import date, datetime

import openpyxl
from django.db.models import Q
from django.http import HttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.exceptions import NotFound
from rest_framework.parsers import MultiPartParser, FormParser
from rest_framework.permissions import IsAuthenticated
from rest_framework.response import Response

from apps.files.drive_service import upload_file, delete_file
from .models import BMRTrackerRecord, BMRTrackerFile
from .serializers import BMRTrackerRecordSerializer, BMRTrackerFileSerializer

TEMPLATE_COLUMNS = ['date', 'client_name', 'product', 'batch_no']

MAX_UPLOAD_SIZE_BYTES = 10 * 1024 * 1024  # 10 MB

_DATE_FORMATS = [
    '%Y-%m-%d', '%d-%m-%Y', '%d/%m/%Y', '%m/%d/%Y',
    '%b %Y', '%B %Y', '%b-%Y', '%B-%Y',
]


def _parse_date(raw):
    if raw is None or raw == '':
        return None
    if isinstance(raw, datetime):
        return raw.date()
    if isinstance(raw, date):
        return raw
    raw = str(raw).strip()
    for fmt in _DATE_FORMATS:
        try:
            return datetime.strptime(raw, fmt).date()
        except ValueError:
            continue
    return None


class BMRTrackerRecordViewSet(viewsets.ModelViewSet):
    serializer_class = BMRTrackerRecordSerializer
    permission_classes = [IsAuthenticated]

    def get_queryset(self):
        qs = BMRTrackerRecord.objects.all()
        q = self.request.query_params.get('q')
        if q:
            qs = qs.filter(
                Q(client_name__icontains=q) | Q(product__icontains=q) |
                Q(batch_no__icontains=q)
            )
        return qs

    def perform_create(self, serializer):
        serializer.save(created_by=self.request.user)

    def destroy(self, request, *args, **kwargs):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete.'}, status=403)
        return super().destroy(request, *args, **kwargs)

    # ------------------------------------------------------------------
    # GET /api/bmr-tracker/upload-template/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['get'], url_path='upload-template')
    def download_template(self, request):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'BMR Tracker'

        header_font = openpyxl.styles.Font(bold=True)
        header_fill = openpyxl.styles.PatternFill('solid', fgColor='FFF3CD')

        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            cell = ws.cell(row=1, column=col_idx, value=col_name)
            cell.font = header_font
            cell.fill = header_fill

        widths = [14, 22, 26, 18]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = width

        hint_font = openpyxl.styles.Font(italic=True, color='888888')
        date_col = TEMPLATE_COLUMNS.index('date') + 1
        ws.cell(row=2, column=date_col, value='YYYY-MM-DD, DD-MM-YYYY, or month + year e.g. Feb 2026').font = hint_font

        sample = {
            'date': '2026-06-01',
            'client_name': 'Vedanshi',
            'product': 'Vitamin C Serum',
            'batch_no': 'B2026001',
        }
        for col_idx, col_name in enumerate(TEMPLATE_COLUMNS, start=1):
            ws.cell(row=3, column=col_idx, value=sample[col_name])

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)

        response = HttpResponse(
            buf.read(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = 'attachment; filename="bmr_tracker_template.xlsx"'
        return response

    # ------------------------------------------------------------------
    # POST /api/bmr-tracker/bulk-upload/
    # ------------------------------------------------------------------
    @action(detail=False, methods=['post'], url_path='bulk-upload')
    def bulk_upload(self, request):
        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file provided. Send an Excel file as "file".'}, status=400)

        try:
            wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
        except Exception:
            return Response({'detail': 'Could not read the file. Please upload a valid .xlsx file.'}, status=400)

        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return Response({'detail': 'The file is empty.'}, status=400)

        header_row = [str(c).strip().lower() if c is not None else '' for c in rows[0]]

        def col(name):
            try:
                return header_row.index(name)
            except ValueError:
                return None

        idx = {name: col(name) for name in TEMPLATE_COLUMNS}

        def cell(row_data, name):
            i = idx[name]
            if i is None or i >= len(row_data):
                return None
            v = row_data[i]
            return str(v).strip() if v is not None else None

        created = []
        to_create = []

        for row_num, row_data in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == '' for v in row_data):
                continue

            date_raw = row_data[idx['date']] if idx['date'] is not None and idx['date'] < len(row_data) else None
            parsed_date = _parse_date(date_raw)

            to_create.append(BMRTrackerRecord(
                date=parsed_date,
                client_name=cell(row_data, 'client_name') or '',
                product=cell(row_data, 'product') or '',
                batch_no=cell(row_data, 'batch_no') or '',
                created_by=request.user,
            ))
            entry = {'row': row_num, 'client_name': cell(row_data, 'client_name') or '—'}
            if date_raw and not parsed_date:
                entry['warning'] = f'Could not parse date "{date_raw}" — left blank'
            created.append(entry)

        BMRTrackerRecord.objects.bulk_create(to_create)
        wb.close()
        return Response({'created': created, 'skipped': []}, status=200)


class BMRTrackerFilesViewSet(viewsets.ViewSet):
    permission_classes = [IsAuthenticated]
    parser_classes = [MultiPartParser, FormParser]

    def list(self, request, bmr_tracker_id=None):
        files = BMRTrackerFile.objects.filter(bmr_tracker_record_id=bmr_tracker_id).select_related('uploaded_by')
        return Response(BMRTrackerFileSerializer(files, many=True).data)

    def create(self, request, bmr_tracker_id=None):
        try:
            record = BMRTrackerRecord.objects.get(pk=bmr_tracker_id)
        except BMRTrackerRecord.DoesNotExist:
            raise NotFound('BMR tracker record not found')

        file_obj = request.FILES.get('file')
        if not file_obj:
            return Response({'detail': 'No file uploaded'}, status=status.HTTP_400_BAD_REQUEST)
        if file_obj.size > MAX_UPLOAD_SIZE_BYTES:
            return Response({'detail': 'File exceeds the 10 MB upload limit.'}, status=status.HTTP_400_BAD_REQUEST)

        try:
            result = upload_file(
                file_bytes=file_obj.read(),
                filename=file_obj.name,
                mimetype=file_obj.content_type or 'application/octet-stream',
                client_name=record.client_name or 'BMR Tracker',
                subfolder='BMR Tracker',
            )
        except Exception as e:
            return Response({'detail': f'Drive upload failed: {e}'}, status=status.HTTP_502_BAD_GATEWAY)

        file_record = BMRTrackerFile.objects.create(
            bmr_tracker_record=record,
            drive_file_id=result['drive_file_id'],
            drive_url=result['drive_url'],
            filename=file_obj.name,
            uploaded_by=request.user,
        )
        return Response(BMRTrackerFileSerializer(file_record).data, status=status.HTTP_201_CREATED)

    def destroy(self, request, pk=None, bmr_tracker_id=None):
        if request.user.role != 'admin':
            return Response({'detail': 'Only admin can delete files.'}, status=status.HTTP_403_FORBIDDEN)
        try:
            file_record = BMRTrackerFile.objects.get(pk=pk, bmr_tracker_record_id=bmr_tracker_id)
        except BMRTrackerFile.DoesNotExist:
            raise NotFound()
        delete_file(file_record.drive_file_id)
        file_record.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)
